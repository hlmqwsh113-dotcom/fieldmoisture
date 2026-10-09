# -*- coding: utf-8 -*-
"""excel_io 回归测试：读 → 模拟键盘录入 → 保存 → 再读，全链路验证。"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from excel_io import load_rows, save_rows, default_output_path
from openpyxl import Workbook, load_workbook

tmp = tempfile.mkdtemp()
path = os.path.join(tmp, "t.xlsx")

# 构造带"干扰"的表：列序打乱 + 已有部分水分值 + 水分列不在末尾
wb = Workbook()
ws = wb.active
ws.append(["排", "田间条码", "品种名称", "水分", "区"])
ws.append([1, "T001", "先玉335", None, 2])
ws.append([1, "T002", "郑单958", 12.5, 3])
ws.append([2, "T003", "京科968", None, 1])
ws.append([None, None, None, None, None])   # 空行应被跳过
ws.append([2, "T004", "登海605", 0.8, 4])
wb.save(path)

rows, missing = load_rows(path)
assert missing == [], "不应缺表头: %s" % missing
assert len(rows) == 4, "空行应被跳过, got %d" % len(rows)
assert rows[0] == {"code": "T001", "pai": "1", "qu": "2",
                   "variety": "先玉335", "moisture": ""}, rows[0]
assert rows[1]["moisture"] == "12.5", rows[1]
assert rows[3]["moisture"] == "0.8", rows[3]
print("PASS: 乱序列读取 / 空行跳过 / 品种名称 / 既有数值保留")

# 旧表头"田间编号"兼容
path_old = os.path.join(tmp, "old.xlsx")
wb2 = Workbook()
ws2 = wb2.active
ws2.append(["田间编号", "排", "区"])
ws2.append(["T201", 1, 1])
wb2.save(path_old)
rows_old, missing_old = load_rows(path_old)
assert rows_old[0]["code"] == "T201"
assert missing_old == [u"品种名称", u"水分"], missing_old
print("PASS: 旧表头 田间编号 兼容")

# 模拟键盘录入
rows[0]["moisture"] = "15.3"
rows[2]["moisture"] = "0"
save_rows(path, rows)

rows2, _ = load_rows(path)
assert rows2[0]["moisture"] == "15.3"
assert rows2[2]["moisture"] == "0"
assert rows2[1]["moisture"] == "12.5"
assert rows2[3]["moisture"] == "0.8"

# 验证写回了正确的单元格（乱序列下 T001 的水分应在 D 列）
wb3 = load_workbook(path)
assert wb3.active["D2"].value == 15.3
assert wb3.active["A3"].value == 1       # 排列未被破坏
assert wb3.active["B2"].value == "T001"
assert wb3.active["C2"].value == "先玉335"
print("PASS: 保存后数值落位正确，其他列不受影响")

# 缺"水分"列的表 → 保存时自动补列
path3 = os.path.join(tmp, "t3.xlsx")
wb4 = Workbook()
ws4 = wb4.active
ws4.append(["田间条码", "品种名称", "排", "区"])
ws4.append(["T101", "先玉335", 1, 1])
wb4.save(path3)
rows3, missing3 = load_rows(path3)
assert missing3 == [u"水分"]
rows3[0]["moisture"] = "7.7"
save_rows(path3, rows3)
rows4, _ = load_rows(path3)
assert rows4[0]["moisture"] == "7.7"
print("PASS: 缺水分列时自动补列")

assert default_output_path(path).endswith("t_已填写.xlsx")
print("PASS: 全部测试通过")
